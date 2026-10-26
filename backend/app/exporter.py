import os
import re
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    SchedulePeriod,
    ScheduleAssignment,
    Employee,
    Shift,
)


def export_period(db: Session, period_id):
    period = db.get(SchedulePeriod, period_id)

    if not period:
        raise ValueError("period not found")

    # 【2026-10-09 修复】与系统各处保持一致：按 sort_order、再按姓名排序。
    # 原来用 order_by(Employee.name)，导出的人员顺序与排班表/统计页不一致。
    shifts = {
        x.id: x
        for x in db.scalars(select(Shift)).all()
    }

    rows = db.scalars(
        select(ScheduleAssignment).where(
            ScheduleAssignment.period_id == period.id
        )
    ).all()

    # 【2026-10-10】导出同样要包含"该周期内有记录的已停用员工"，
    # 否则停用某人后，其历史周期的导出会整行消失（数据还在库里）。
    data_ids = {r.employee_id for r in rows}

    employees = [
        x
        for x in db.scalars(
            select(Employee).order_by(Employee.sort_order, Employee.name)
        ).all()
        if x.active or x.id in data_ids
    ]

    data = defaultdict(dict)

    for row in rows:
        data[row.employee_id][row.work_date] = (
            shifts[row.shift_id].code
            if row.shift_id
            else ""
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "排班表"

    ws["A1"] = period.name
    ws["A1"].font = Font(bold=True, size=14)

    ws.cell(2, 1).value = "Engineer"
    ws.cell(2, 2).value = "Email"

    dates = []
    d = period.start_date

    while d <= period.end_date:
        dates.append(d)
        d = d.fromordinal(d.toordinal() + 1)

    for i, d in enumerate(dates, start=3):
        ws.cell(1, i).value = d.strftime("%m/%d")
        ws.cell(2, i).value = d.strftime("%a")

    for r, e in enumerate(employees, start=3):
        ws.cell(r, 1).value = e.name
        ws.cell(r, 2).value = e.email or ""

        for c, d in enumerate(dates, start=3):
            ws.cell(r, c).value = data[e.id].get(d, "")

    for column in ws.columns:
        for cell in column:
            cell.alignment = Alignment(horizontal="center")

    os.makedirs("/data/exports", exist_ok=True)

    # 【修复】period.name 直接拼路径，含 / 等字符会报错或穿越，先清洗
    safe_name = re.sub(r"[^\w\u4e00-\u9fff\-]+", "_", period.name).strip("_") or "period"
    path = f"/data/exports/{safe_name}.xlsx"

    wb.save(path)

    return path
