"""生成端到端验证用的合成 Excel（结构模拟真实 CSE 班表）。

在 api 容器内运行（容器里有 openpyxl）：
    docker compose run --rm api python /data/make_test_excel.py

输出：/data/uploads/CSE_2026年班表--排班.xlsx
- 12 名员工（6 人有 N/Y 夜班权限），31 天列（对应 26→25 周期）
- 含 1 行汇总行（应被跳过）+ 1 个 O4 单元格（应产生 warning）
"""
import random
from openpyxl import Workbook

random.seed(42)

NIGHT_OK = ["张伟", "李娜", "王强", "刘洋", "陈静", "赵磊"]      # 6 人可上夜班
ALL = ["张伟", "李娜", "王强", "刘洋", "陈静", "赵磊",
       "孙敏", "周杰", "吴磊", "郑爽", "冯军", "褚丽"]          # 12 人

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"

ws.cell(1, 1).value = "CSE 2026年8月份排班表"

# 第 2 行星期占位（导入器只作显示用，不读取）
ws.cell(2, 1).value = "Tier"

for col in range(5, 36):
    ws.cell(2, col).value = "周" + "一二三四五六日"[(col - 5 + 2) % 7]

for i, name in enumerate(ALL):
    r = 3 + i
    ws.cell(r, 1).value = "T1(XA)" if i % 2 == 0 else "T1(CQ)"
    ws.cell(r, 3).value = name
    ws.cell(r, 4).value = f"{name.lower()}@telstra-pbs.cn"

    for col in range(5, 36):
        roll = random.random()
        if name in NIGHT_OK:
            if roll < 0.14:
                code = "N"
            elif roll < 0.18:
                code = "Y"
            elif roll < 0.33:
                code = "D"
            elif roll < 0.45:
                code = "O2"
            elif roll < 0.57:
                code = "O3"
            elif roll < 0.64:
                code = "C"
            else:
                code = "O1"
        else:
            if roll < 0.30:
                code = "O2"
            elif roll < 0.55:
                code = "O3"
            elif roll < 0.68:
                code = "D"
            elif roll < 0.75:
                code = "C"
            else:
                code = "O1"
        ws.cell(r, col).value = code

    # 给第 1 名员工塞一个历史 O4，验证报警逻辑
    if i == 0:
        ws.cell(r, 35).value = "O4"

# 汇总行（无姓名，应被跳过）
sr = 3 + len(ALL)
ws.cell(sr, 1).value = "总数"
for col in range(5, 36):
    ws.cell(sr, col).value = 7

wb.save("/data/uploads/CSE_2026年班表--排班.xlsx")
print("test excel written: /data/uploads/CSE_2026年班表--排班.xlsx")
